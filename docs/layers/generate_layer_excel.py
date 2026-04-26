# -*- coding: utf-8 -*-
"""从 doc/Layers 下 v1.1 Markdown 生成可编辑 Excel（需 openpyxl）。"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CATALOG_MD = ROOT / "001-图层完整清单-v1.md"
MAPPING_MD = ROOT / "002-旧到新图层映射-v1.md"
OUT_XLSX = ROOT / "HyCAD图层-v1.1.xlsx"


def strip_cell(s: str) -> str:
    s = s.strip()
    if s.startswith("`") and s.endswith("`"):
        s = s[1:-1]
    return s


def is_sep_row(cells: list[str]) -> bool:
    if not cells:
        return True
    first = cells[0].strip()
    if first.startswith("---"):
        return True
    if all(re.fullmatch(r"-+", c.strip().replace(" ", "")) for c in cells if c.strip()):
        return True
    return False


def parse_catalog_rows(md_path: Path) -> list[dict]:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    major_code = ""
    major_name = ""
    subgroup = ""
    rows: list[dict] = []
    # 仅解析「## 1. `00` …」「## 2. `01` …」「## 3. `02` …」正文，避免 §6.3 等附表误入
    catalog_active = False

    major_re = re.compile(r"^##\s+\d+\.\s+`(\d{2})`\s+(.+)$")
    section_end_re = re.compile(r"^##\s+[4-9]\.\s+")
    sub_re = re.compile(r"^###\s+[\d.]+\s+`([^`]+)`")

    for line in lines:
        if section_end_re.match(line):
            catalog_active = False
        m = major_re.match(line)
        if m:
            major_code, major_name = m.group(1), m.group(2).strip()
            catalog_active = major_code in ("00", "01", "02")
            continue
        m = sub_re.search(line)
        if m:
            subgroup = m.group(1)
            continue
        if not catalog_active or not line.startswith("|"):
            continue
        cells = [strip_cell(c) for c in line.split("|")[1:-1]]
        if len(cells) < 2 or is_sep_row(cells):
            continue
        if cells[0] == "新层名" or cells[0] == "字段":
            continue
        if cells[0].startswith(("00-hy-", "01-hy-", "02-hy-")):
            rows.append(
                {
                    "major_code": major_code,
                    "major_name": major_name,
                    "subgroup": subgroup,
                    "cells": cells,
                }
            )
    return rows


def refine_catalog_sheet() -> list[list]:
    """修正 parse：聚类表 9 列且第 8 列为「归属（合并前）」+ 第 9 列 XData。"""
    parsed = parse_catalog_rows(CATALOG_MD)
    fixed: list[list] = []
    for r in parsed:
        c = r["cells"]
        n = len(c)
        xdata = ""
        if n == 9:
            name, sid, color, lt, lw, lock, plot, owner, xdata = c
        elif n == 8:
            name, sid, color, lt, lw, lock, plot, owner = c
        else:
            name = c[0]
            sid, color, lt, lw, lock, plot, owner = (c + [""] * 8)[1:8]
            xdata = c[8] if n > 8 else ""

        is_tpl = "是" if ("{" in name and "}" in name) or str(color) == "动态" else "否"
        fixed.append(
            [
                r["major_code"],
                r["major_name"],
                r["subgroup"],
                name,
                sid,
                color,
                lt,
                lw,
                lock,
                plot,
                owner,
                xdata,
                is_tpl,
            ]
        )
    return fixed


def parse_mapping_rows(md_path: Path) -> tuple[list[str], list[list]]:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    headers = ["章节", "旧层名", "新层名", "映射类型", "XData_KIND", "来源说明"]
    rows: list[list] = []
    section = ""
    section_re = re.compile(r"^##\s+(\d+)\.\s+`(\d{2})`")

    for line in lines:
        m = section_re.match(line)
        if m:
            section = f"{m.group(2)}-{m.group(1)}"
            continue
        if not line.startswith("|"):
            continue
        cells = [strip_cell(c) for c in line.split("|")[1:-1]]
        if len(cells) < 3 or is_sep_row(cells):
            continue
        if cells[0] == "旧层名" or cells[0] == "#":
            continue
        if cells[0].startswith("项目"):
            continue
        # 映射表：旧 | 新 | 类型 | XData | 来源
        if len(cells) >= 5 and cells[1].startswith(("00-hy-", "01-hy-", "02-hy-")):
            old, new, typ, xk, src = cells[0], cells[1], cells[2], cells[3], cells[4]
            rows.append([section, old, new, typ, xk, src])
        elif len(cells) >= 4 and cells[1].startswith(("00-hy-", "01-hy-", "02-hy-")):
            rows.append([section, cells[0], cells[1], cells[2], cells[3], cells[4] if len(cells) > 4 else ""])
    return headers, rows


def parse_xdata_merge_table(md_path: Path) -> tuple[list[str], list[list]]:
    text = md_path.read_text(encoding="utf-8")
    # 提取「## 4. 合并重构说明」后的第一个表格
    idx = text.find("## 4. 合并重构说明")
    if idx < 0:
        return ["#", "合并组", "旧层数", "新层", "XData注册簇", "KIND枚举"], []
    sub = text[idx:]
    lines = sub.splitlines()
    headers = ["#", "合并组", "旧层数", "新层", "XData注册簇", "KIND枚举"]
    rows = []
    in_table = False
    for line in lines:
        if line.startswith("|") and "合并组" in line:
            in_table = True
            continue
        if not in_table:
            continue
        if not line.startswith("|"):
            break
        cells = [strip_cell(c) for c in line.split("|")[1:-1]]
        if is_sep_row(cells) or cells[0] == "#":
            continue
        if cells[0].isdigit() or (cells[0] and cells[0][0].isdigit()):
            rows.append(cells[:6] if len(cells) >= 6 else cells + [""] * (6 - len(cells)))
        else:
            break
    return headers, rows


def parse_aci_table(md_path: Path) -> tuple[list[str], list[list]]:
    text = md_path.read_text(encoding="utf-8")
    idx = text.find("## 4. ACI 颜色索引参考")
    if idx < 0:
        return ["ACI", "颜色", "典型用途"], []
    sub = text[idx : idx + 2500]
    headers = ["ACI", "颜色", "典型用途"]
    rows = []
    for line in sub.splitlines():
        if not line.startswith("|"):
            continue
        cells = [strip_cell(c) for c in line.split("|")[1:-1]]
        if len(cells) < 3 or cells[0] == "ACI" or is_sep_row(cells):
            continue
        if cells[0].isdigit() or cells[0] in ("200", "252", "253"):
            rows.append(cells[:3])
    return headers, rows


def autosize(ws):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 10
        for row in range(1, min(ws.max_row, 200) + 1):
            v = ws.cell(row=row, column=col).value
            if v is not None:
                max_len = max(max_len, min(len(str(v)), 60))
        ws.column_dimensions[letter].width = max_len + 2


def style_header(ws, ncols: int):
    fill = PatternFill("solid", fgColor="FF4472C4")
    font = Font(bold=True, color="FFFFFFFF")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main():
    wb = Workbook()

    # --- 说明 ---
    ws0 = wb.active
    ws0.title = "说明"
    ws0["A1"] = "HyCAD 图层表 v1.1（由 Markdown 自动生成，可自由修改）"
    ws0["A2"] = f"源文件：{CATALOG_MD.name} / {MAPPING_MD.name}"
    ws0["A3"] = (
        "工作表：图层完整清单 = 新规范层（与 001 §1–§3 表一致）；"
        "旧到新映射 = 002 迁移表；XData合并 = §4 注册簇一览；ACI颜色 = 色卡。"
    )
    ws0["A4"] = (
        "行数：清单 92 行（含 2 个动态模板）；映射 110 行。"
        "重新生成：在 doc/Layers 下执行 python generate_layer_excel.py"
    )
    ws0["A5"] = "修改后若需同步回仓库，请更新对应 .md 或后续 LayerCatalogFactory 数据源。"
    for r in range(1, 6):
        ws0.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
    ws0.column_dimensions["A"].width = 100

    # --- 图层完整清单 ---
    parsed = parse_catalog_rows(CATALOG_MD)
    headers = [
        "大类号",
        "大类名",
        "子组",
        "新层名",
        "语义ID",
        "ACI色",
        "线型",
        "线宽mm",
        "锁定",
        "打印",
        "归属",
        "XData说明",
        "是否模板",
    ]
    data = refine_catalog_sheet()
    ws1 = wb.create_sheet("图层完整清单", 1)
    ws1.append(headers)
    for row in data:
        ws1.append(row)
    style_header(ws1, len(headers))
    ws1.freeze_panes = "A2"
    autosize(ws1)

    # --- 旧到新映射 ---
    mh, mrows = parse_mapping_rows(MAPPING_MD)
    ws2 = wb.create_sheet("旧到新映射", 2)
    ws2.append(mh)
    for row in mrows:
        ws2.append(row)
    style_header(ws2, len(mh))
    ws2.freeze_panes = "A2"
    autosize(ws2)

    # --- XData 合并一览 ---
    xh, xrows = parse_xdata_merge_table(MAPPING_MD)
    ws3 = wb.create_sheet("XData合并", 3)
    ws3.append(xh)
    for row in xrows:
        ws3.append(row)
    style_header(ws3, len(xh))
    ws3.freeze_panes = "A2"
    autosize(ws3)

    # --- ACI 颜色 ---
    ah, arows = parse_aci_table(CATALOG_MD)
    ws4 = wb.create_sheet("ACI颜色", 4)
    ws4.append(ah)
    for row in arows:
        ws4.append(row)
    style_header(ws4, len(ah))
    ws4.freeze_panes = "A2"
    autosize(ws4)

    wb.save(OUT_XLSX)
    print(
        f"Wrote {OUT_XLSX} (catalog {len(data)} rows, mapping {len(mrows)}, xdata {len(xrows)})"
    )


if __name__ == "__main__":
    main()
